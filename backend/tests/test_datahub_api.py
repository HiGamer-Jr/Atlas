from pathlib import Path
from uuid import uuid4

import pytest
from cryptography.fernet import Fernet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.datahub.connectors.excel import ExcelConnector
from app.datahub.models import DataHubRecord
from app.datahub.raw_store import RawStore
from tests.test_datahub_preview import policy_configured as _policy_fixture
from tests.test_datahub_preview import preview_case as _case_fixture
from tests.test_datahub_preview import product, upload

policy_configured = _policy_fixture
preview_case = _case_fixture


@pytest.fixture
def enabled_api(app, preview_case):
    settings = app.state.settings.model_copy(
        update={
            "datahub_enabled": True,
            "datahub_raw_root": preview_case[2] / "api-raw",
            "datahub_raw_key": __import__("pydantic").SecretStr(
                Fernet.generate_key().decode()
            ),
        }
    )
    app.state.settings = settings
    app.state.datahub_raw_store = RawStore(settings)
    app.state.datahub_connector = ExcelConnector(
        Path(__file__).resolve().parents[2] / "frontend/src/assets/hiatlas-light.png"
    )
    return preview_case


def test_default_official_logo_resolves(settings, tmp_path):
    from app.main import create_app

    enabled = settings.model_copy(
        update={
            "datahub_enabled": True,
            "datahub_raw_root": tmp_path / "raw",
            "datahub_raw_key": __import__("pydantic").SecretStr(
                Fernet.generate_key().decode()
            ),
        }
    )
    app = create_app(enabled)
    app.state.database_engine.dispose()


def test_template_api_and_download(member, enabled_api):
    headers = enabled_api[1]
    response = member.get("/api/datahub/templates", headers=headers)
    assert response.status_code == 200, response.text
    templates = response.json()["items"]
    assert len(templates) == 8
    assert not next(t for t in templates if t["id"] == "FINANCEIRO")["available"]
    download = member.post(
        "/api/datahub/templates/COMPRADOR_NACIONAL/download", json={}, headers=headers
    )
    assert download.status_code == 200
    assert download.content.startswith(b"PK")
    assert download.headers["cache-control"] == "no-store"


def test_upload_preview_confirm_and_history(member, enabled_api, db_runtime):
    from io import BytesIO

    from openpyxl import load_workbook

    headers = enabled_api[1]
    download = member.post(
        "/api/datahub/templates/COMPRADOR_NACIONAL/download", json={}, headers=headers
    )
    assert download.status_code == 200, download.text
    workbook = load_workbook(BytesIO(download.content))
    sheet = workbook["Produtos"]
    for index, value in enumerate(["000123", "Produto sintético", "UN", "", "SIM"], 1):
        sheet.cell(13, index, value)
    content = BytesIO()
    workbook.save(content)
    response = member.post(
        "/api/datahub/imports",
        files={
            "file": (
                "synthetic.xlsx",
                content.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=headers,
    )
    assert response.status_code == 201, response.text
    preview = response.json()
    assert preview["status"] == "READY_FOR_CONFIRMATION"
    with Session(db_runtime) as db:
        assert db.scalar(select(DataHubRecord)) is None
    rows = member.get(f"/api/datahub/imports/{preview['id']}/rows", headers=headers)
    assert (
        rows.status_code == 200
        and rows.json()["items"][0]["payload"]["codigo"] == "000123"
    )
    command = {"expected_version": preview["version"], "idempotency_key": str(uuid4())}
    first = member.post(
        f"/api/datahub/imports/{preview['id']}/confirm", json=command, headers=headers
    )
    assert first.status_code == 200, first.text
    assert first.json()["status"] == "COMMITTED"
    assert (
        member.post(
            f"/api/datahub/imports/{preview['id']}/confirm",
            json=command,
            headers=headers,
        ).json()
        == first.json()
    )
    assert member.get("/api/datahub/imports", headers=headers).json()["total"] == 1
    exported = member.post(
        "/api/datahub/exports",
        json={"template_id": "COMPRADOR_NACIONAL"},
        headers=headers,
    )
    assert exported.status_code == 200 and exported.content.startswith(b"PK")


def test_confirm_api_rejects_client_payload(member, enabled_api, db_runtime):
    preview = upload(db_runtime, enabled_api, (product(),))
    response = member.post(
        f"/api/datahub/imports/{preview.id}/confirm",
        headers=enabled_api[1],
        json={
            "expected_version": preview.version,
            "idempotency_key": str(uuid4()),
            "before": {"SECRET": "synthetic"},
        },
    )
    assert response.status_code == 422, response.text
    assert "SECRET" not in response.text


@pytest.mark.parametrize("bad", ["origin", "csrf"])
def test_upload_csrf_and_origin(member, enabled_api, bad):
    headers = {
        **enabled_api[1],
        **(
            {"Origin": "https://foreign.example.test"}
            if bad == "origin"
            else {"X-CSRF-Token": "wrong"}
        ),
    }
    response = member.post(
        "/api/datahub/imports",
        files={"file": ("synthetic.xlsx", b"invalid")},
        headers=headers,
    )
    assert response.status_code == 403, response.text


def test_foreign_import_is_neutral(member, enabled_api):
    response = member.get(f"/api/datahub/imports/{uuid4()}", headers=enabled_api[1])
    assert response.status_code == 404, response.text


def test_support_and_admin_role_have_no_implicit_access(support, admin, enabled_api):
    from tests.helpers import select_context

    for client in [support, admin]:
        context = select_context(client, enabled_api[0]["contract_a"])
        response = client.get("/api/datahub/templates", headers=context)
        assert response.status_code == 403, response.text


@pytest.mark.parametrize(
    "contract,role,email",
    [
        ("contract_a2", "role_a2", "member@example.test"),
        ("contract_b", "role_b", "other@example.test"),
    ],
)
def test_valid_foreign_ids_are_hidden_across_contracts(
    enabled_api, db_runtime, new_client, contract, role, email
):
    from sqlalchemy import text

    from tests.helpers import select_context
    from tests.identity_helpers import login

    ids = enabled_api[0]
    preview = upload(db_runtime, enabled_api, (product(),))
    tenant = ids["tenant_b" if contract == "contract_b" else "tenant_a"]
    with db_runtime.begin() as c:
        c.execute(
            text(
                "INSERT INTO contract_modules (tenant_id,contract_id,code,contracted,active) VALUES (:t,:c,'DATAHUB',true,true)"
            ),
            {"t": tenant, "c": ids[contract]},
        )
        for cap in [
            "datahub.read",
            "datahub.import",
            "datahub.products.read",
            "datahub.products.import",
        ]:
            c.execute(
                text(
                    "INSERT INTO tenant_role_permissions (tenant_id,contract_id,role_id,capability) VALUES (:t,:c,:r,:cap)"
                ),
                {"t": tenant, "c": ids[contract], "r": ids[role], "cap": cap},
            )
    client = new_client()
    login(client, email)
    context = select_context(client, ids[contract])
    assert client.get("/api/datahub/imports", headers=context).json()["total"] == 0
    for path in [
        str(preview.id),
        str(preview.id) + "/rows",
        str(preview.id) + "/issues",
    ]:
        assert (
            client.get("/api/datahub/imports/" + path, headers=context).status_code
            == 404
        )
    assert (
        client.post(
            f"/api/datahub/imports/{preview.id}/confirm",
            headers=context,
            json={"expected_version": preview.version, "idempotency_key": str(uuid4())},
        ).status_code
        == 404
    )


def test_new_http_session_reads_history_but_cannot_confirm(
    enabled_api, db_runtime, new_client
):
    from tests.helpers import select_context
    from tests.identity_helpers import login

    preview = upload(db_runtime, enabled_api, (product(),))
    other = new_client()
    login(other, "member@example.test")
    context = select_context(other, enabled_api[0]["contract_a"])
    assert (
        other.get(f"/api/datahub/imports/{preview.id}", headers=context).status_code
        == 200
    )
    response = other.post(
        f"/api/datahub/imports/{preview.id}/confirm",
        headers=context,
        json={"expected_version": preview.version, "idempotency_key": str(uuid4())},
    )
    assert response.status_code == 404, response.text


def test_upload_budget_and_invalid_package(member, enabled_api, app):
    app.state.settings = app.state.settings.model_copy(
        update={
            "datahub_limits": app.state.settings.datahub_limits.model_copy(
                update={"max_upload_bytes": 16}
            )
        }
    )
    assert (
        member.post(
            "/api/datahub/imports",
            files={"file": ("synthetic.xlsx", b"x" * 20)},
            headers=enabled_api[1],
        ).status_code
        == 413
    )
    assert (
        member.post(
            "/api/datahub/imports",
            files={"file": ("synthetic.xlsx", b"invalid")},
            headers=enabled_api[1],
        ).status_code
        == 422
    )
