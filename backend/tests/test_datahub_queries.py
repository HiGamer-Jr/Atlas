from importlib.util import find_spec
from uuid import uuid4

import pytest
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.core.errors import ApiError
from tests.test_datahub_confirmation import confirm
from tests.test_datahub_policy import context_values
from tests.test_datahub_preview import policy_configured as _policy_fixture
from tests.test_datahub_preview import preview_case as _case_fixture
from tests.test_datahub_preview import product, upload

policy_configured = _policy_fixture
preview_case = _case_fixture


def query_call(engine, case, name, *args, settings=None):
    assert find_spec("app.datahub.queries"), "Authorized history/export missing"
    from app.datahub import queries

    with Session(engine) as db, db.begin():
        principal, scope = context_values(db, case[1])
        if settings:
            db.info["settings"] = settings
        return getattr(queries, name)(db, principal, scope, *args)


def test_history_and_projections_have_no_raw_reference(preview_case, db_runtime):
    preview = upload(db_runtime, preview_case, (product(),))
    assert find_spec("app.datahub.queries"), "Authorized history/export missing"
    from app.datahub.schemas import ImportQuery

    page = query_call(db_runtime, preview_case, "list_imports", ImportQuery())
    assert page.total == 1
    detail = query_call(db_runtime, preview_case, "get_import", preview.id)
    assert detail.filename == "synthetic.xlsx"
    rows = query_call(db_runtime, preview_case, "list_rows", preview.id, ImportQuery())
    assert rows.total == 1 and rows.items[0].payload["codigo"] == "000123"
    assert "raw_reference" not in detail.model_dump_json()
    assert "fingerprint" not in rows.model_dump_json()


@pytest.mark.parametrize("foreign", [True, False])
def test_foreign_or_forbidden_import_is_neutral(preview_case, db_runtime, foreign):
    preview = upload(db_runtime, preview_case, (product(),))
    if not foreign:
        with db_runtime.begin() as c:
            c.execute(
                text(
                    "UPDATE tenant_role_permissions SET active=false WHERE role_id=:r AND capability='datahub.products.read'"
                ),
                {"r": preview_case[0]["role_basic"]},
            )
    with pytest.raises(ApiError) as error:
        query_call(
            db_runtime, preview_case, "get_import", uuid4() if foreign else preview.id
        )
    assert error.value.status in (403, 404)


def test_export_limit_never_silently_truncates(preview_case, db_runtime):
    preview = upload(db_runtime, preview_case, (product("P-ONE"), product("P-TWO")))
    confirm(db_runtime, preview_case, preview)
    assert find_spec("app.datahub.queries"), "Authorized export missing"
    from app.core.config import Settings
    from app.datahub.schemas import ExportInput

    with pytest.raises(ApiError) as error:
        query_call(
            db_runtime,
            preview_case,
            "export_workbook",
            ExportInput(template_id="COMPRADOR_NACIONAL"),
            settings=Settings(
                datahub_enabled=True,
                datahub_raw_root=preview_case[2] / "raw",
                datahub_raw_key=__import__("cryptography.fernet", fromlist=["Fernet"])
                .Fernet.generate_key()
                .decode(),
                datahub_limits={"max_export_rows": 1},
            ),
        )
    assert error.value.status == 413


def test_export_is_official_and_preserves_text(preview_case, db_runtime):
    preview = upload(
        db_runtime, preview_case, (product(description="=synthetic literal"),)
    )
    confirm(db_runtime, preview_case, preview)
    assert find_spec("app.datahub.queries"), "Authorized export missing"
    from io import BytesIO

    from openpyxl import load_workbook

    from app.datahub.schemas import ExportInput

    download = query_call(
        db_runtime,
        preview_case,
        "export_workbook",
        ExportInput(template_id="COMPRADOR_NACIONAL"),
    )
    workbook = load_workbook(BytesIO(download.content))
    assert workbook["Produtos"]["A13"].value == "000123"
    assert workbook["Produtos"]["B13"].data_type == "s"
    assert "LEIA-ME" in workbook.sheetnames


def test_history_rechecks_current_unit_scope(preview_case, db_runtime):
    from app.datahub.schemas import DemandPayload, ImportQuery
    from app.datahub.types import NormalizedRow

    demand = NormalizedRow(
        "DEMANDS",
        1,
        "Demandas",
        13,
        DemandPayload(
            codigo="D-SYN",
            produto_codigo="000123",
            unidade_codigo="DH-UNIT",
            quantidade="1",
            data_necessidade="2026-10-20",
            modalidade="NACIONAL",
            prioridade="NORMAL",
        ),
    )
    preview = upload(db_runtime, preview_case, (product(), demand))
    assert (
        query_call(db_runtime, preview_case, "list_imports", ImportQuery()).total == 1
    )
    with db_runtime.begin() as c:
        c.execute(
            text(
                "UPDATE membership_unit_scopes SET active=false WHERE membership_id=:m"
            ),
            {"m": preview_case[0]["member_a"]},
        )
    assert (
        query_call(db_runtime, preview_case, "list_imports", ImportQuery()).total == 0
    )
    for name, args in [
        ("get_import", ()),
        ("list_rows", (ImportQuery(),)),
        ("list_issues", (ImportQuery(),)),
    ]:
        with pytest.raises(ApiError) as error:
            query_call(db_runtime, preview_case, name, preview.id, *args)
        assert error.value.status == 404


def test_history_pagination_and_sanitized_filename(preview_case, db_runtime):
    from app.datahub.schemas import ImportQuery

    one = upload(
        db_runtime, preview_case, (product(),), filename="../../synthetic.xlsx"
    )
    upload(db_runtime, preview_case, (product("000124"),))
    page = query_call(
        db_runtime, preview_case, "list_imports", ImportQuery(page_size=1)
    )
    assert page.total == 2 and len(page.items) == 1
    assert (
        query_call(db_runtime, preview_case, "get_import", one.id).filename
        == "synthetic.xlsx"
    )
    assert (
        query_call(
            db_runtime, preview_case, "list_imports", ImportQuery(page=2, page_size=1)
        )
        .items[0]
        .id
        == one.id
    )


def test_export_template_filters_modality_before_limit(preview_case, db_runtime):
    from io import BytesIO

    from openpyxl import load_workbook

    from app.datahub.schemas import DemandPayload, ExportInput
    from app.datahub.types import NormalizedRow
    from tests.test_datahub_policy import invoke

    demands = tuple(
        NormalizedRow(
            "DEMANDS",
            1,
            "Demandas",
            13,
            DemandPayload(
                codigo="D-" + mode,
                produto_codigo="000123",
                unidade_codigo="DH-UNIT",
                quantidade="1",
                data_necessidade="2026-10-20",
                modalidade=mode,
                prioridade="NORMAL",
            ),
        )
        for mode in ["NACIONAL", "INTERNACIONAL"]
    )
    with db_runtime.begin() as c:
        c.execute(
            text(
                "INSERT INTO tenant_role_permissions (tenant_id,contract_id,role_id,capability) VALUES (:t,:c,:r,'datahub.demands.export')"
            ),
            {
                "t": preview_case[0]["tenant_a"],
                "c": preview_case[0]["contract_a"],
                "r": preview_case[0]["role_basic"],
            },
        )
    preview = upload(
        db_runtime,
        preview_case,
        (product(), *demands),
        selection=invoke(db_runtime, preview_case[1], "COORDENACAO"),
    )
    confirm(db_runtime, preview_case, preview)
    detail = query_call(db_runtime, preview_case, "get_import", preview.id)
    for template, mode in [
        ("COMPRADOR_NACIONAL", "NACIONAL"),
        ("COMPRADOR_INTERNACIONAL", "INTERNACIONAL"),
    ]:
        result = query_call(
            db_runtime,
            preview_case,
            "export_workbook",
            ExportInput(template_id=template),
        )
        workbook = load_workbook(BytesIO(result.content))
        sheet = workbook["Demandas"]
        assert sheet["F13"].value == mode
        assert sheet["A14"].value is None

    assert detail.unit_scope == ["DH-UNIT"]
