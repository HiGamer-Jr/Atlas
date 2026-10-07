from datetime import UTC, date, datetime
from decimal import Decimal
from importlib.util import find_spec
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

import pytest

from app.datahub.catalog import DATASETS, TEMPLATES
from app.datahub.schemas import StockPayload
from app.datahub.types import AuthorizedSelection, AuthorizedUnit


def connector():
    assert find_spec("app.datahub.connectors.excel"), "Secure Excel connector missing"
    from app.datahub.connectors.excel import ExcelConnector

    return ExcelConnector(
        Path(__file__).resolve().parents[2] / "frontend/src/assets/hiatlas-light.png"
    )


def official(template="COMPRADOR_NACIONAL", rows=()):
    excel = connector()
    from app.datahub.types import WorkbookContext

    definition = TEMPLATES[template]
    unit = AuthorizedUnit(uuid4(), "UNIT-SYN", "Unidade sintética", 1)
    selection = AuthorizedSelection(
        template,
        1,
        definition.datasets,
        "Perfil sintético",
        (unit,),
        "Empresa sintética",
        "SYN-001",
        "TEST",
    )
    context = WorkbookContext(
        "Empresa sintética",
        "SYN-001",
        "Perfil sintético",
        "UNIT-SYN",
        datetime(2026, 10, 7, tzinfo=UTC),
    )
    return excel.generate(selection, context, rows)


def parsed(content):
    from app.datahub.types import WorkbookLimits

    return connector().parse(content, WorkbookLimits())


def edit(content, sheet, cell, value):
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(content))
    wb[sheet][cell] = value
    result = BytesIO()
    wb.save(result)
    wb.close()
    return result.getvalue()


@pytest.mark.parametrize("template", tuple(TEMPLATES))
def test_official_layout_and_readme(template):
    content = official(template)
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(content))
    assert wb["LEIA-ME"].sheet_state == "visible"
    assert wb["_HIATLAS_META"].sheet_state == "hidden"
    readme = " ".join(str(c.value) for row in wb["LEIA-ME"] for c in row if c.value)
    for term in (
        "não concede autorização",
        "obrigatórios",
        "não renomear",
        "_HIATLAS_META",
        "versão",
    ):
        assert term.lower() in readme.lower()
    for code in TEMPLATES[template].datasets:
        ws = wb[DATASETS[code].sheet]
        assert ws.freeze_panes == "A13"
        assert (
            tuple(c.value for c in ws[12])[: len(DATASETS[code].fields)]
            == DATASETS[code].fields
        )
        assert len(ws._images) == 1
        assert ws.tables
        assert ws.data_validations.dataValidation
        values = " ".join(
            str(c.value)
            for row in ws.iter_rows(min_row=1, max_row=11)
            for c in row
            if c.value
        )
        assert "HiAtlas — Supply Chain Intelligence" in values
        assert "Um novo horizonte para o seu negócio" in values
        assert "Empresa sintética" in values and "SYN-001" in values
    assert not any(c.data_type == "f" for ws in wb for row in ws for c in row)
    wb.close()


def test_codes_preserve_leading_zeroes_and_numeric_codes_are_rejected():
    content = official()
    content = edit(content, "Produtos", "A13", "000123")
    content = edit(content, "Produtos", "B13", "Produto sintético")
    result = parsed(content)
    assert result.rows[0].values["codigo"] == "000123"
    bad = parsed(edit(content, "Produtos", "A13", 123))
    assert any(
        i.code == "TEXT_CODE_REQUIRED" and i.source_row == 13 and i.column == "codigo"
        for i in bad.issues
    )


def test_formula_cache_is_not_a_value():
    content = edit(official(), "Produtos", "A13", "=1+1")
    source = ZipFile(BytesIO(content))
    out = BytesIO()
    with ZipFile(out, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename.startswith("xl/worksheets/"):
                data = data.replace(b"<f>1+1</f><v></v>", b"<f>1+1</f><v>123</v>")
            target.writestr(info, data)
    source.close()
    result = parsed(out.getvalue())
    assert any(i.code == "FORMULA_NOT_ALLOWED" for i in result.issues)
    assert result.rows[0].values["codigo"] is None


@pytest.mark.parametrize(
    "value",
    [
        "=1+1",
        '  =HYPERLINK("https://example.test")',
        "\t=1+1",
        "+cmd",
        "@SUM(A1)",
        "-text",
    ],
)
def test_export_literal_prefixes_and_negative_decimal(value):
    excel = connector()
    from openpyxl import Workbook, load_workbook

    wb = Workbook()
    ws = wb.active
    excel.write_cell(ws["A1"], value)
    excel.write_cell(ws["A2"], Decimal("-12.2500"))
    out = BytesIO()
    wb.save(out)
    result = load_workbook(BytesIO(out.getvalue()))
    assert result.active["A1"].data_type == "s"
    assert result.active["A1"].value == value
    assert result.active["A1"].hyperlink is None
    assert result.active["A2"].data_type == "n" and result.active["A2"].value == -12.25


def test_large_decimal_roundtrip_preserves_exact_value():
    connector()
    from app.datahub.types import NormalizedRow

    payload = StockPayload(
        produto_codigo="000123",
        unidade_codigo="UNIT-SYN",
        quantidade_disponivel="9999999999999999.1234",
        data_referencia=date(2026, 10, 7),
    )
    row = NormalizedRow("STOCK_POSITIONS", 1, "Estoque", 13, payload)
    result = parsed(official("CENTRO_DISTRIBUICAO", (row,)))
    assert result.rows[0].values["quantidade_disponivel"] == "9999999999999999.1234"


@pytest.mark.parametrize(
    "threat",
    [
        "traversal",
        "macro",
        "external_link",
        "xml_entity",
        "unknown_sheet",
        "duplicate_column",
        "changed_manifest",
    ],
)
def test_untrusted_package_and_structure_fail_closed(threat):
    content = official()
    excel = connector()
    from app.datahub.connectors.base import WorkbookRejected
    from app.datahub.types import WorkbookLimits

    if threat == "duplicate_column":
        content = edit(content, "Produtos", "B12", "codigo")
    elif threat == "changed_manifest":
        content = edit(content, "_HIATLAS_META", "A1", "tampered")
    elif threat == "unknown_sheet":
        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(content))
        wb.create_sheet("Unknown")
        out = BytesIO()
        wb.save(out)
        content = out.getvalue()
    else:
        out = BytesIO()
        with (
            ZipFile(BytesIO(content)) as source,
            ZipFile(out, "w", ZIP_DEFLATED) as target,
        ):
            for info in source.infolist():
                data = source.read(info.filename)
                if threat == "xml_entity" and info.filename == "xl/workbook.xml":
                    data = (
                        b'<!DOCTYPE x [<!ENTITY payload SYSTEM "file:///private">]>'
                        + data
                    )
                target.writestr(info, data)
            if threat == "traversal":
                target.writestr("../escape", b"x")
            if threat == "macro":
                target.writestr("xl/vbaProject.bin", b"x")
            if threat == "external_link":
                target.writestr("xl/externalLinks/externalLink1.xml", b"<x/>")
        content = out.getvalue()
    with pytest.raises(WorkbookRejected):
        excel.parse(content, WorkbookLimits())


def test_configurable_package_limits_are_enforced():
    content = official()
    from app.datahub.connectors.base import WorkbookRejected
    from app.datahub.types import WorkbookLimits

    with pytest.raises(WorkbookRejected) as error:
        connector().parse(content, WorkbookLimits(max_upload_bytes=32))
    assert error.value.code == "FILE_TOO_LARGE"
    with pytest.raises(WorkbookRejected):
        connector().parse(content, WorkbookLimits(max_uncompressed_bytes=32))


def test_missing_first_cell_remains_a_row_for_field_validation():
    from openpyxl import load_workbook

    wb = load_workbook(
        BytesIO(edit(official(), "Produtos", "B13", "Produto sintético"))
    )
    wb["Produtos"]["A13"].style = "Normal"
    output = BytesIO()
    wb.save(output)
    wb.close()
    result = parsed(output.getvalue())
    assert result.rows[0].source_row == 13
    assert result.rows[0].values["codigo"] is None


@pytest.mark.parametrize("cell", ["Z12", "Z13"])
def test_nonadjacent_unknown_columns_are_rejected(cell):
    content = edit(official(), "Produtos", cell, "unknown")
    from app.datahub.connectors.base import WorkbookRejected

    with pytest.raises(WorkbookRejected):
        parsed(content)


def test_generated_metadata_identifies_profile_context_and_utc_generation():
    content = official()
    import json

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(content))
    metadata = json.loads(wb["_HIATLAS_META"]["A2"].value)
    assert metadata["information"]["profile"] == "Perfil sintético"
    assert metadata["information"]["company"] == "Empresa sintética"
    assert metadata["information"]["generated_at"].startswith("2026-10-07")
    wb.close()


def test_manifest_version_does_not_accept_boolean_or_float():
    content = official()
    import json

    from openpyxl import load_workbook

    from app.datahub.connectors.base import WorkbookRejected

    for value in (True, 1.0):
        wb = load_workbook(BytesIO(content))
        manifest = json.loads(wb["_HIATLAS_META"]["A2"].value)
        manifest["version"] = value
        wb["_HIATLAS_META"]["A2"] = json.dumps(manifest)
        out = BytesIO()
        wb.save(out)
        wb.close()
        with pytest.raises(WorkbookRejected):
            parsed(out.getvalue())


def test_generator_uses_configured_data_validation_range():
    from openpyxl import load_workbook

    from app.datahub.connectors.excel import ExcelConnector
    from app.datahub.types import WorkbookContext, WorkbookLimits

    excel = ExcelConnector(
        Path(__file__).resolve().parents[2] / "frontend/src/assets/hiatlas-light.png",
        limits=WorkbookLimits(max_rows=20),
    )
    selection = AuthorizedSelection(
        "COMPRADOR_NACIONAL",
        1,
        ("PRODUCTS",),
        "Perfil",
        (),
        "Empresa",
        "SYN-001",
        "TEST",
    )
    content = excel.generate(
        selection,
        WorkbookContext(
            "Empresa", "SYN-001", "Perfil", "Sem unidades", datetime.now(UTC)
        ),
        (),
    )
    wb = load_workbook(BytesIO(content))
    ranges = [str(v.sqref) for v in wb["Produtos"].data_validations.dataValidation]
    assert ranges and all(r.endswith("32") for r in ranges)


def test_text_serializer_never_silently_truncates():
    from openpyxl import Workbook

    from app.datahub.connectors.base import WorkbookRejected

    wb = Workbook()
    with pytest.raises(WorkbookRejected):
        connector().write_cell(wb.active["A1"], "x" * 32768)


def test_row_limit_applies_to_data_not_hidden_enum_ranges():
    content = edit(official(), "Produtos", "A13", "000123")
    from app.datahub.types import WorkbookLimits

    result = connector().parse(content, WorkbookLimits(max_rows=1))
    assert len(result.rows) == 1


def test_modality_validation_matches_selected_template():
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(official("COMPRADOR_NACIONAL")))
    dv = next(
        v
        for v in wb["Demandas"].data_validations.dataValidation
        if str(v.sqref).startswith("F13")
    )
    import re

    col, first, last = re.search(
        r"\$([A-Z]+)\$(\d+):\$[A-Z]+\$(\d+)", dv.formula1
    ).groups()
    choices = [
        wb["_HIATLAS_META"][f"{col}{r}"].value for r in range(int(first), int(last) + 1)
    ]
    assert choices == ["NACIONAL"]


def test_numeric_xml_roundtrip_avoids_float_conversion():
    from app.datahub.types import NormalizedRow

    payload = StockPayload(
        produto_codigo="000123",
        unidade_codigo="UNIT-SYN",
        quantidade_disponivel="1.0000",
        data_referencia=date(2026, 10, 7),
    )
    content = official(
        "CENTRO_DISTRIBUICAO",
        (NormalizedRow("STOCK_POSITIONS", 1, "Estoque", 13, payload),),
    )
    out = BytesIO()
    from xml.etree import ElementTree as ET

    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with ZipFile(BytesIO(content)) as source, ZipFile(out, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename.startswith("xl/worksheets/"):
                root = ET.fromstring(data)
                cell = root.find(".//s:c[@r='C13']/s:v", ns)
                if cell is not None:
                    cell.text = "1234567890123456.1234"
                    data = ET.tostring(root)
            target.writestr(info, data)
    result = parsed(out.getvalue())
    assert result.rows[0].values["quantidade_disponivel"] == Decimal(
        "1234567890123456.1234"
    )


@pytest.mark.parametrize(
    "limit", [{"max_entries": 1}, {"max_sheets": 1}, {"max_cells": 1}]
)
def test_package_entry_sheet_and_cell_limits(limit):
    content = official()
    from app.datahub.connectors.base import WorkbookRejected
    from app.datahub.types import WorkbookLimits

    with pytest.raises(WorkbookRejected):
        connector().parse(content, WorkbookLimits(**limit))


def test_generator_never_emits_unassigned_unit_rows():
    from app.datahub.connectors.base import WorkbookRejected
    from app.datahub.types import NormalizedRow

    payload = StockPayload(
        produto_codigo="000123",
        unidade_codigo="FOREIGN",
        quantidade_disponivel="1.0000",
        data_referencia=date(2026, 10, 7),
    )
    with pytest.raises(WorkbookRejected):
        official(
            "CENTRO_DISTRIBUICAO",
            (NormalizedRow("STOCK_POSITIONS", 1, "Estoque", 13, payload),),
        )


@pytest.mark.parametrize("coordinate", ["!BAD", "A0", "A-1", "XFE13"])
def test_malformed_coordinate_rejected_without_library_payload(coordinate):
    from xml.etree import ElementTree as ET

    from app.datahub.connectors.base import WorkbookRejected

    out = BytesIO()
    content = official()
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with ZipFile(BytesIO(content)) as source, ZipFile(out, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet2.xml":
                root = ET.fromstring(data)
                root.find(".//s:sheetData/s:row/s:c", ns).set("r", coordinate)
                data = ET.tostring(root)
            target.writestr(info, data)
    with pytest.raises(WorkbookRejected) as error:
        parsed(out.getvalue())
    assert str(error.value) == "PACKAGE_INVALID"


def test_corrupt_deflate_is_sanitized_package_error():
    import struct

    from app.datahub.connectors.base import WorkbookRejected

    content = bytearray(official())
    with ZipFile(BytesIO(content)) as source:
        entry = source.getinfo("xl/workbook.xml")
        name_length, extra_length = struct.unpack_from(
            "<HH", content, entry.header_offset + 26
        )
        start = entry.header_offset + 30 + name_length + extra_length
    content[start : start + 4] = b"\xff" * 4
    with pytest.raises(WorkbookRejected) as error:
        parsed(bytes(content))
    assert str(error.value) == "PACKAGE_INVALID"
