"""Attested synthetic E2E factory. Never imported by the production application."""

import argparse
import json
import os
import sys
from pathlib import Path
from uuid import uuid4

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
from sqlalchemy import select, text
from sqlalchemy.orm import Session

from app.core.config import Settings
from app.datahub.models import DataHubRecord
from app.db.base import Base
from app.main import create_app
from tests.database_harness import open_test_engines, validate_test_urls
from tests.helpers import seed_scope
from tests.identity_helpers import seed_user


def guard():
    owner, runtime = validate_test_urls(
        os.environ.get("TEST_DATABASE_OWNER_URL", ""),
        os.environ.get("TEST_DATABASE_RUNTIME_URL", ""),
    )
    if (
        os.environ.get("HIATLAS_DATAHUB_E2E") != "1"
        or (owner.host, owner.port, owner.database)
        != ("127.0.0.1", 55493, "hiatlas_datahub_test")
        or os.environ.get("PUBLIC_ORIGIN") != "https://localhost:5188"
    ):
        raise ValueError("Explicit owned disposable E2E configuration required")
    return owner, runtime


def engines():
    guard()
    return open_test_engines()


def clean():
    owner, runtime = engines()
    try:
        with owner.begin() as c:
            c.execute(
                text(
                    "TRUNCATE "
                    + ", ".join('"' + t.name + '"' for t in Base.metadata.sorted_tables)
                    + " CASCADE"
                )
            )
    finally:
        owner.dispose()
        runtime.dispose()


def seed():
    owner, runtime = engines()
    try:
        with runtime.connect() as c:
            if c.scalar(text("SELECT count(*) FROM users")):
                raise ValueError("Empty owned synthetic database required")
        users = {
            "admin_user": seed_user(runtime, "admin@example.test", "PLATFORM_ADMIN"),
            "support_user": seed_user(
                runtime, "support@example.test", "PLATFORM_SUPPORT"
            ),
            "member_user": seed_user(runtime, "member@example.test"),
        }
        ids = seed_scope(runtime, users)
        from app.datahub.catalog import DATASETS
        from app.identity.passwords import hash_password

        password = os.environ["HIATLAS_E2E_PASSWORD"]
        if len(password) < 20:
            raise ValueError("Protected synthetic credential required")
        with runtime.begin() as c:
            c.execute(
                text(
                    "UPDATE users SET password_hash=:hash,display_name='Pessoa sintética'"
                ),
                {"hash": hash_password(password)},
            )
            c.execute(
                text(
                    "INSERT INTO memberships (id,tenant_id,contract_id,user_id,role_id) VALUES (:id,:t,:c,:u,:r)"
                ),
                {
                    "id": uuid4(),
                    "t": ids["tenant_a"],
                    "c": ids["contract_a"],
                    "u": users["admin_user"],
                    "r": ids["role_basic"],
                },
            )
            for contract, tenant, role, member in [
                ("contract_a", "tenant_a", "role_basic", "member_a"),
                ("contract_a2", "tenant_a", "role_a2", "member_a2"),
                ("contract_b", "tenant_b", "role_b", "member_b"),
            ]:
                for module in ["DATAHUB", "PROCUREMENT", "COMEX", "INVENTORY"]:
                    c.execute(
                        text(
                            "INSERT INTO contract_modules (tenant_id,contract_id,code,contracted,active) VALUES (:t,:c,:code,true,true)"
                        ),
                        {"t": ids[tenant], "c": ids[contract], "code": module},
                    )
                caps = [
                    "datahub.read",
                    "datahub.template.download",
                    "datahub.import",
                    "datahub.export",
                ] + [
                    f"datahub.{d.lower()}.{op}"
                    for d, definition in DATASETS.items()
                    if not definition.sensitive
                    for op in ["read", "import", "export"]
                ]
                for cap in caps:
                    c.execute(
                        text(
                            "INSERT INTO tenant_role_permissions (tenant_id,contract_id,role_id,capability) VALUES (:t,:c,:r,:cap)"
                        ),
                        {
                            "t": ids[tenant],
                            "c": ids[contract],
                            "r": ids[role],
                            "cap": cap,
                        },
                    )
                node = uuid4()
                ids["unit_" + contract] = node
                c.execute(
                    text(
                        "INSERT INTO organization_nodes (id,tenant_id,contract_id,kind,name,code) VALUES (:id,:t,:c,'STORE','Unidade sintética','UNIT-SYN')"
                    ),
                    {"id": node, "t": ids[tenant], "c": ids[contract]},
                )
                c.execute(
                    text(
                        "INSERT INTO membership_unit_scopes (tenant_id,contract_id,membership_id,node_id) VALUES (:t,:c,:m,:n)"
                    ),
                    {"t": ids[tenant], "c": ids[contract], "m": ids[member], "n": node},
                )
                if contract != "contract_a":
                    c.execute(
                        text(
                            "INSERT INTO memberships (id,tenant_id,contract_id,user_id,role_id) VALUES (:id,:t,:c,:u,:r)"
                        ),
                        {
                            "id": uuid4(),
                            "t": ids[tenant],
                            "c": ids[contract],
                            "u": users["admin_user"],
                            "r": ids[role],
                        },
                    )
                c.execute(
                    text(
                        "INSERT INTO membership_unit_scopes (tenant_id,contract_id,membership_id,node_id) SELECT tenant_id,contract_id,id,:n FROM memberships WHERE user_id=:u AND contract_id=:c"
                    ),
                    {"n": node, "u": users["admin_user"], "c": ids[contract]},
                )
        return {k: str(v) for k, v in ids.items()}
    finally:
        owner.dispose()
        runtime.dispose()


def create_browser_app():
    guard()
    owner, runtime = engines()
    owner.dispose()
    runtime.dispose()
    settings = Settings(
        environment="test",
        database_url=os.environ["TEST_DATABASE_RUNTIME_URL"],
        public_origin=os.environ["PUBLIC_ORIGIN"],
        datahub_enabled=True,
        datahub_raw_root=Path(os.environ["HIATLAS_E2E_STATE"]) / "raw",
        datahub_raw_key=os.environ["HIATLAS_DATAHUB_KEY"],
    )
    app = create_app(settings)
    from app.identity.email_transport import FakeEmailTransport

    app.state.email_transport = FakeEmailTransport()
    return app


def fill(source, variant):
    root = Path(os.environ["HIATLAS_E2E_STATE"]).resolve()
    source = Path(source).resolve()
    if not source.is_relative_to(root):
        raise ValueError("Private synthetic XLSX path required")
    from openpyxl import load_workbook

    workbook = load_workbook(source)
    values = {
        "Produtos": ["000123", "Produto sintético", "UN", "Sintética", "SIM"],
        "Parceiros": ["PART-SYN", "Parceiro sintético", "FORNECEDOR", "BR", "SIM"],
        "Demandas": [
            "DEM-SYN",
            "000123",
            "UNIT-SYN",
            12.5,
            "2026-10-20",
            "NACIONAL",
            "NORMAL",
        ],
        "Estoque": ["000123", "UNIT-SYN", 10, "2026-10-07"],
        "COMEX": [
            "COMEX-SYN",
            "PART-SYN",
            "USD",
            "FOB",
            "2026-10-25",
            "PLANEJADO",
        ],
    }
    for sheet, row in values.items():
        if sheet not in workbook.sheetnames:
            continue
        for index, value in enumerate(row, 1):
            workbook[sheet].cell(13, index, value)
    if variant == "invalid":
        workbook["Produtos"]["B13"] = "=1+1"
    elif variant == "changed":
        workbook["Produtos"]["B13"] = "Produto sintético divergente"
    target = root / (variant + ".xlsx")
    workbook.save(target)
    return str(target)


def inspect():
    owner, runtime = engines()
    try:
        with Session(runtime) as db:
            from app.audit.models import AuditEvent

            return {
                "records": len(list(db.scalars(select(DataHubRecord.id)))),
                "committed_events": len(
                    list(
                        db.scalars(
                            select(AuditEvent.id).where(
                                AuditEvent.action == "datahub.import.committed"
                            )
                        )
                    )
                ),
            }
    finally:
        owner.dispose()
        runtime.dispose()


def revoke(kind):
    owner, runtime = engines()
    try:
        with runtime.begin() as c:
            if kind == "permission":
                c.execute(
                    text(
                        "UPDATE tenant_role_permissions SET active=false WHERE capability='datahub.products.import' AND contract_id=(SELECT id FROM contracts WHERE code='CTR-A')"
                    )
                )
            elif kind == "module":
                c.execute(
                    text(
                        "UPDATE contract_modules SET active=false WHERE code='DATAHUB' AND contract_id=(SELECT id FROM contracts WHERE code='CTR-A')"
                    )
                )
    finally:
        owner.dispose()
        runtime.dispose()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "action", choices=["clean", "seed", "fill", "inspect", "revoke"]
    )
    parser.add_argument("--source")
    parser.add_argument("--variant", default="valid")
    args = parser.parse_args()
    try:
        guard()
        result = (
            clean()
            if args.action == "clean"
            else seed()
            if args.action == "seed"
            else fill(args.source, args.variant)
            if args.action == "fill"
            else inspect()
            if args.action == "inspect"
            else revoke(args.variant)
        )
        print(json.dumps(result))
    except Exception:  # noqa: BLE001 -- sensitive operational errors are never printed
        print("Synthetic E2E fixture failed; diagnostics redacted", file=sys.stderr)
        raise SystemExit(1) from None
